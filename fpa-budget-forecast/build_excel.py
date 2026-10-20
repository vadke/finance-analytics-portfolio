"""Build the Excel FP&A model (formulas, not pasted results) from the analysis outputs."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
d = json.loads((HERE / "output" / "dashboard_data.json").read_text())
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E79")
B = Font(name=F, bold=True); N = Font(name=F); BLUE = Font(name=F, color="0000FF"); GREEN = Font(name=F, color="008000")
YEL = PatternFill("solid", fgColor="FFFF00"); TOP = Border(top=Side(style="thin"))
USD = '$#,##0;($#,##0);"-"'; USDK = '$#,##0,;($#,##0,);"-"'; PCT = '0.0%;(0.0%);"-"'
LINES = ["Revenue", "Cost of Goods Sold", "Store Operations", "Marketing", "Supply Chain & Logistics", "Technology", "General & Administrative"]
OPEX = LINES[2:]

wb = Workbook()
wb.properties.creator = "Vinit Adke"
wb.properties.lastModifiedBy = "Vinit Adke"
wb.properties.title = "Summit Outfitters FP&A Model"
# ---------------- README ----------------
ws = wb.active; ws.title = "README"
notes = ["Summit Outfitters FP&A Model (simulated data)",
 "",
 "Purpose: FY2026 budget vs. actual variance analysis and a 12-month rolling forecast with scenarios.",
 "Data: simulated monthly P&L for a fictional specialty retailer; margins calibrated to public big-box retail peers.",
 "Actuals run through Sep 2026. Oct 2026 onward is forecast.",
 "",
 "Sheets",
 "  Assumptions: scenario selector and forecast drivers. Blue = input, yellow = key assumption you can change.",
 "  Data: monthly budget and actual by P&L line (source: data/pl_monthly.csv, summed across regions and channels).",
 "  Budget vs Actual: FY2026 by month with variance $ and %, YTD and full-year latest estimate.",
 "  Forecast: Oct 2026 - Sep 2027 P&L driven by the selected scenario.",
 "",
 "Sign convention: revenue positive, costs negative, so variance = actual - budget is favorable when positive.",
 "Base revenue forecast: log-linear trend + month-of-year model fitted per region x channel in analysis.py "
 f"(backtest MAPE {d['backtest']['mape_pct']}% on Jan-Sep 2026)."]
for i, t in enumerate(notes, 1):
    ws.cell(i, 1, t).font = Font(name=F, bold=(i in (1, 7)), size=14 if i == 1 else 10)
ws.column_dimensions["A"].width = 120

# ---------------- Assumptions ----------------
a = wb.create_sheet("Assumptions")
a["A1"] = "Forecast assumptions"; a["A1"].font = Font(name=F, bold=True, size=13)
a["A3"] = "Selected scenario"; a["B3"] = "Base"; a["B3"].font = BLUE; a["B3"].fill = YEL
dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False); a.add_data_validation(dv); dv.add("B3")
a["C3"] = "Choose Base, Upside or Downside"
for c, t in zip("ABC", ["Scenario", "Revenue multiplier", "Gross margin adj. (pts)"]):
    a[f"{c}5"] = t; a[f"{c}5"].font = H; a[f"{c}5"].fill = HF
sc = d["forecast"]["assumptions"]["scenarios"]
for i, (k, v) in enumerate(sc.items(), 6):
    a[f"A{i}"] = k
    a[f"B{i}"] = v["revenue_multiplier"]; a[f"B{i}"].number_format = "0.00x"
    a[f"C{i}"] = v["gross_margin_adj"]; a[f"C{i}"].number_format = PCT
    for c in "BC": a[f"{c}{i}"].font = BLUE; a[f"{c}{i}"].fill = YEL
a["A10"] = "Active revenue multiplier"; a["B10"] = "=INDEX($B$6:$B$8,MATCH($B$3,$A$6:$A$8,0))"; a["B10"].number_format = "0.00x"
a["A11"] = "Active gross margin adj."; a["B11"] = "=INDEX($C$6:$C$8,MATCH($B$3,$A$6:$A$8,0))"; a["B11"].number_format = PCT
a["A12"] = "Gross margin, trailing 12 months"; a["B12"] = d["forecast"]["assumptions"]["gross_margin_t12"]; a["B12"].number_format = PCT; a["B12"].font = BLUE
a["C12"] = "Source: actuals Oct 2025 - Sep 2026 (analysis.py)"
for c, t in zip("ABC", ["Department", "Fixed cost / month ($)", "Variable rate (% of revenue)"]):
    a[f"{c}14"] = t; a[f"{c}14"].font = H; a[f"{c}14"].fill = HF
om = d["forecast"]["assumptions"]["opex_model"]
for i, dept in enumerate(OPEX, 15):
    a[f"A{i}"] = dept
    a[f"B{i}"] = om[dept]["fixed_monthly"]; a[f"B{i}"].number_format = USD
    a[f"C{i}"] = om[dept]["variable_rate"]; a[f"C{i}"].number_format = '0.00%'
    for c in "BC": a[f"{c}{i}"].font = BLUE
a["A21"] = "Fixed and variable splits are OLS fits of monthly department cost on revenue, trailing 12 months (analysis.py)."
for c, w in zip("ABC", [36, 24, 30]): a.column_dimensions[c].width = w

# ---------------- Data ----------------
ds = wb.create_sheet("Data")
for c, t in enumerate(["Month", "Scenario", "Line item", "Amount ($)"], 1):
    ds.cell(1, c, t).font = H; ds.cell(1, c).fill = HF
r = 2
bva = [x for x in d["bva_monthly"] if x["fiscal_year"] == "2026" and x["line_item"] in LINES]
for x in bva:
    for scen, val in (("Budget", x["budget"]), ("Actual", x["actual"])):
        if val is None: continue
        ds.cell(r, 1, x["month"]); ds.cell(r, 2, scen); ds.cell(r, 3, x["line_item"])
        ds.cell(r, 4, val).number_format = USD; ds.cell(r, 4).font = BLUE; r += 1
last_data = r - 1
fm = d["forecast"]["months"]; base_rev = d["forecast"]["scenarios"]["Base"]["revenue"]
ds.cell(1, 6, "Forecast month").font = H; ds.cell(1, 6).fill = HF
ds.cell(1, 7, "Base revenue forecast ($)").font = H; ds.cell(1, 7).fill = HF
for i, (m, v) in enumerate(zip(fm, base_rev), 2):
    ds.cell(i, 6, m); ds.cell(i, 7, v).number_format = USD; ds.cell(i, 7).font = BLUE
ds.cell(15, 6, "Source: analysis.py model output (sum of region x channel models)")
for c, w in zip("ABCDEFG", [10, 10, 26, 16, 3, 16, 24]): ds.column_dimensions[c].width = w
RNG = lambda col: f"Data!${col}$2:${col}${last_data}"

# ---------------- Budget vs Actual ----------------
bv = wb.create_sheet("Budget vs Actual")
months26 = [f"2026-{m:02d}" for m in range(1, 13)]
bv["A1"] = "FY2026 Budget vs. Actual ($000s)"; bv["A1"].font = Font(name=F, bold=True, size=13)
bv["A2"] = "Actuals through Sep 2026. Variance = Actual - Budget; positive is favorable."
hdr = ["Line item", "Type"] + [m for m in months26] + ["YTD Budget", "YTD Actual", "YTD Var $", "YTD Var %", "FY Budget", "FY Latest Est.", "FY Var $"]
for c, t in enumerate(hdr, 1):
    cell = bv.cell(4, c, t); cell.font = H; cell.fill = HF; cell.alignment = Alignment(horizontal="center", wrap_text=True)
order = LINES[:2] + ["Gross Profit"] + OPEX + ["Operating Income"]
row = 5; rows_of = {}
for line in order:
    for typ in ("Budget", "Actual", "Variance $"):
        bv.cell(row, 1, line if typ == "Budget" else ""); bv.cell(row, 2, typ)
        rows_of[(line, typ)] = row; row += 1
    row += 0
def ref(line, typ, col): return f"{L(col)}{rows_of[(line, typ)]}"
for line in order:
    for ci, m in enumerate(months26, 3):
        col = L(ci)
        for typ in ("Budget", "Actual"):
            rr = rows_of[(line, typ)]
            if line == "Gross Profit":
                f = f"={ref('Revenue', typ, ci)}+{ref('Cost of Goods Sold', typ, ci)}"
            elif line == "Operating Income":
                f = f"={ref('Gross Profit', typ, ci)}+" + "+".join(ref(o, typ, ci) for o in OPEX)
            else:
                f = f'=SUMIFS({RNG("D")},{RNG("A")},"{m}",{RNG("B")},"{typ}",{RNG("C")},"{line}")'
            if typ == "Actual" and m > "2026-09":
                f = ""
            bv.cell(rr, ci, f).number_format = USDK
        rv = rows_of[(line, "Variance $")]
        bv.cell(rv, ci, f'=IF({col}{rows_of[(line,"Actual")]}="","",{col}{rows_of[(line,"Actual")]}-{col}{rows_of[(line,"Budget")]})').number_format = USDK
    # YTD and FY columns: O..U (15..21)
    rb, ra, rv = rows_of[(line, "Budget")], rows_of[(line, "Actual")], rows_of[(line, "Variance $")]
    bv.cell(rb, 15, f"=SUM(C{rb}:K{rb})"); bv.cell(rb, 16, f"=SUM(C{ra}:K{ra})")
    bv.cell(rb, 17, f"=P{rb}-O{rb}"); bv.cell(rb, 18, f"=IF(O{rb}=0,0,Q{rb}/ABS(O{rb}))")
    bv.cell(rb, 19, f"=SUM(C{rb}:N{rb})")
    for c in (15, 16, 17, 19): bv.cell(rb, c).number_format = USDK
    bv.cell(rb, 18).number_format = PCT
    for c in range(1, 22): bv.cell(rv, c).border = Border(bottom=Side(style="hair"))
    for c in range(3, 22): bv.cell(rv, c).font = Font(name=F, italic=True)
# FY latest estimate = YTD actual + Oct-Dec forecast (from Forecast sheet, columns B..D)
fc_row = {"Revenue": 5, "Cost of Goods Sold": 6, "Gross Profit": 7, "Operating Income": 14}
for i, o in enumerate(OPEX): fc_row[o] = 8 + i
for line in order:
    rb = rows_of[(line, "Budget")]
    fr = fc_row[line]
    bv.cell(rb, 20, f"=P{rb}+SUM(Forecast!B{fr}:D{fr})").font = GREEN
    bv.cell(rb, 21, f"=T{rb}-S{rb}")
    for c in (20, 21): bv.cell(rb, c).number_format = USDK
for line in ("Gross Profit", "Operating Income"):
    for typ in ("Budget", "Actual", "Variance $"):
        for c in range(1, 22): bv.cell(rows_of[(line, typ)], c).font = Font(name=F, bold=True, italic=(typ == "Variance $"))
bv.column_dimensions["A"].width = 26; bv.column_dimensions["B"].width = 11
for c in range(3, 22): bv.column_dimensions[L(c)].width = 11
bv.freeze_panes = "C5"

# ---------------- Forecast ----------------
fs = wb.create_sheet("Forecast")
fs["A1"] = "Rolling forecast, Oct 2026 - Sep 2027 ($000s)"; fs["A1"].font = Font(name=F, bold=True, size=13)
fs["A2"] = '="Scenario: "&Assumptions!B3'; fs["A2"].font = GREEN
fs.cell(4, 1, "Line item").font = H; fs.cell(4, 1).fill = HF
for ci, m in enumerate(fm, 2):
    c = fs.cell(4, ci, m); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center")
c = fs.cell(4, 14, "12-mo total"); c.font = H; c.fill = HF
labels = ["Revenue", "Cost of Goods Sold", "Gross Profit"] + OPEX + ["Total Opex", "Operating Income", "Operating margin %"]
for i, t in enumerate(labels, 5): fs.cell(i, 1, t)
# rows: 5 rev, 6 cogs, 7 gp, 8-12 opex, 13 total opex, 14 OI, 15 margin
for ci in range(2, 14):
    col = L(ci); dr = ci  # Data!G2..G13
    fs[f"{col}5"] = f"=Data!$G${dr}*Assumptions!$B$10"
    fs[f"{col}6"] = f"=-{col}5*(1-(Assumptions!$B$12+Assumptions!$B$11))"
    fs[f"{col}7"] = f"={col}5+{col}6"
    for j in range(5):
        fs[f"{col}{8+j}"] = f"=-(Assumptions!$B${15+j}+Assumptions!$C${15+j}*{col}$5)"
    fs[f"{col}13"] = f"=SUM({col}8:{col}12)"
    fs[f"{col}14"] = f"={col}7+{col}13"
    fs[f"{col}15"] = f"=IF({col}5=0,0,{col}14/{col}5)"
    for r_ in range(5, 15): fs[f"{col}{r_}"].number_format = USDK
    fs[f"{col}15"].number_format = PCT
for r_ in range(5, 15):
    fs[f"N{r_}"] = f"=SUM(B{r_}:M{r_})"; fs[f"N{r_}"].number_format = USDK
fs["N15"] = "=IF(N5=0,0,N14/N5)"; fs["N15"].number_format = PCT
for r_ in (7, 14):
    for c in range(1, 15): fs.cell(r_, c).font = B; fs.cell(r_, c).border = TOP
fs.column_dimensions["A"].width = 26
for c in range(2, 15): fs.column_dimensions[L(c)].width = 12
fs.freeze_panes = "B5"

for sh in wb.worksheets:
    for row_ in sh.iter_rows():
        for cell in row_:
            if cell.font.name != F:
                cell.font = Font(name=F, bold=cell.font.bold, italic=cell.font.italic, color=cell.font.color, size=cell.font.size)
out = HERE / "output" / "Summit_Outfitters_FPA_Model.xlsx"
wb.save(out); print(out)
